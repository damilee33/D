"""Small, explicit compatibility shims for the pinned local environment."""

import sys


def patch_numpy_legacy_aliases():
    """Restore aliases needed by openpyxl 3.0.5 under NumPy >= 1.24.

    The shim is intentionally local and must run before importing openpyxl.
    It does not change numeric behavior: each alias maps to the corresponding
    Python scalar type.
    """

    import numpy as np

    aliases = {
        "bool": bool,
        "float": float,
        "int": int,
        "object": object,
    }
    for name, value in aliases.items():
        if name not in np.__dict__:
            setattr(np, name, value)


def import_openpyxl_load_workbook():
    """Import openpyxl without loading its unused optional pandas integration.

    The project reads cells directly through openpyxl and never exchanges pandas
    objects with it.  The local pandas/numexpr pair is incompatible, so exposing
    pandas to openpyxl would trigger an irrelevant optional-dependency warning.
    """

    patch_numpy_legacy_aliases()
    sentinel = object()
    previous_pandas = sys.modules.get("pandas", sentinel)
    if previous_pandas is sentinel:
        sys.modules["pandas"] = None
    try:
        from openpyxl import load_workbook
    finally:
        if previous_pandas is sentinel:
            sys.modules.pop("pandas", None)
        else:
            sys.modules["pandas"] = previous_pandas
    return load_workbook
